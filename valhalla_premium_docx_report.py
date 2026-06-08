from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import base64
import os
import tempfile

try:
    from valhalla_strategy_engine import generate_dynamic_tax_strategy
except Exception:
    generate_dynamic_tax_strategy = None

try:
    from valhalla_roi_engine import generate_roi_analysis
except Exception:
    generate_roi_analysis = None

BRAND_RED = "981E26"
BRAND_DARK = "242424"
BRAND_BLACK = "18181B"
BRAND_GOLD = "B7892B"
DEEP_GOLD = "8A651D"
INK = "1F2933"
TEXT_GRAY = "5C6670"
LIGHT_GRAY = "F3F5F7"
LINE_GRAY = "D8DEE5"
LIGHT_RED = "F8ECEE"
LIGHT_GOLD = "FFF8E6"
SOFT_SLATE = "EEF2F6"
WHITE = "FFFFFF"
LOGO_ASSET = os.path.join(os.path.dirname(__file__), "assets", "valhalla_gold_logo_report.b64")


def _money(value):
    try:
        amount = float(value)
        if amount < 0:
            return f"(${abs(amount):,.0f})"
        return f"${amount:,.0f}"
    except Exception:
        return "$0"


def _num(value, default=0):
    try:
        return float(value)
    except Exception:
        return float(default)


def _pct(numerator, denominator):
    denominator = _num(denominator)
    if not denominator:
        return "0.0%"
    return f"{(_num(numerator) / denominator * 100):.1f}%"


def _rate(value):
    try:
        return f"{float(value) * 100:.1f}%"
    except Exception:
        return "0.0%"


def _safe(value, fallback="Not provided"):
    if value is None or value == "":
        return fallback
    return str(value)


def _limit(text, length=420):
    text = _safe(text, "")
    if len(text) <= length:
        return text
    return text[: length - 3].rstrip() + "..."


def _filing_key(status):
    status = str(status or "").strip().lower().replace("-", " ").replace("_", " ")
    if status in ("mfj", "married filing jointly", "married joint", "joint"):
        return "mfj"
    if status in ("hoh", "head of household"):
        return "hoh"
    if status in ("mfs", "married filing separately", "separate"):
        return "mfs"
    return "single"


def _ordinary_brackets(status):
    key = _filing_key(status)
    if key == "mfj":
        return [(23850, 0.10), (96950, 0.12), (206700, 0.22), (394600, 0.24), (501050, 0.32), (751600, 0.35), (10**12, 0.37)]
    if key == "hoh":
        return [(17000, 0.10), (64850, 0.12), (103350, 0.22), (197300, 0.24), (250500, 0.32), (626350, 0.35), (10**12, 0.37)]
    if key == "mfs":
        return [(11925, 0.10), (48475, 0.12), (103350, 0.22), (197300, 0.24), (250525, 0.32), (375800, 0.35), (10**12, 0.37)]
    return [(11925, 0.10), (48475, 0.12), (103350, 0.22), (197300, 0.24), (250525, 0.32), (626350, 0.35), (10**12, 0.37)]


def _ordinary_tax(taxable_income, status):
    taxable = max(0, _num(taxable_income))
    total = 0
    last = 0
    for top, rate in _ordinary_brackets(status):
        if taxable <= last:
            break
        amount = min(taxable, top) - last
        total += amount * rate
        last = top
    return total


def _marginal_rate(taxable_income, status):
    taxable = max(0, _num(taxable_income))
    for top, rate in _ordinary_brackets(status):
        if taxable <= top:
            return rate
    return 0.37


def _ltcg_thresholds(status):
    key = _filing_key(status)
    if key == "mfj":
        return 96700, 600050
    if key == "hoh":
        return 64750, 566700
    if key == "mfs":
        return 48350, 300000
    return 48350, 533400


def _estimate_ltcg_tax(ordinary_base, preferred_income, status):
    ordinary_base = max(0, _num(ordinary_base))
    remaining = max(0, _num(preferred_income))
    zero_top, fifteen_top = _ltcg_thresholds(status)
    zero_room = max(0, zero_top - ordinary_base)
    at_zero = min(remaining, zero_room)
    remaining -= at_zero
    fifteen_room = max(0, fifteen_top - max(ordinary_base + at_zero, zero_top))
    at_fifteen = min(remaining, fifteen_room)
    remaining -= at_fifteen
    at_twenty = max(0, remaining)
    tax = at_fifteen * 0.15 + at_twenty * 0.20
    return at_zero, at_fifteen, at_twenty, tax


def _set_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_border(cell, color=LINE_GRAY, size="4"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        element = borders.find(qn("w:" + edge))
        if element is None:
            element = OxmlElement("w:" + edge)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def _clear_border(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        element = borders.find(qn("w:" + edge))
        if element is None:
            element = OxmlElement("w:" + edge)
            borders.append(element)
        element.set(qn("w:val"), "nil")


def _cell_margins(cell, top=120, start=140, bottom=120, end=140):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + name))
        if node is None:
            node = OxmlElement("w:" + name)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _format_cell(cell, fill=None, color=INK, bold=False, size=8.8, align=None, border=LINE_GRAY):
    if fill:
        _set_shading(cell, fill)
    if border is None:
        _clear_border(cell)
    else:
        _set_border(cell, border)
    _cell_margins(cell)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.05
        if align is not None:
            paragraph.alignment = align
        for run in paragraph.runs:
            run.font.name = "Aptos"
            run.font.size = Pt(size)
            run.bold = bold
            run.font.color.rgb = RGBColor.from_string(color)


def _paragraph(paragraph, size=9.4, color=INK, bold=False, italic=False, before=0, after=5, align=None):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.08
    if align is not None:
        paragraph.alignment = align
    for run in paragraph.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(size)
        run.font.color.rgb = RGBColor.from_string(color)
        if bold:
            run.bold = True
        if italic:
            run.italic = True


def _set_widths(table, widths):
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths):
            if idx < len(row.cells):
                row.cells[idx].width = Inches(width)


def _add_cell_text(cell, text, size=8.8, color=INK, bold=False, after=0, align=None):
    p = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
    p.text = ""
    run = p.add_run(str(text))
    _paragraph(p, size=size, color=color, bold=bold, after=after, align=align)
    return p


def _page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def _logo_temp_path():
    if not os.path.exists(LOGO_ASSET):
        return None
    try:
        with open(LOGO_ASSET, "r", encoding="utf-8") as logo_file:
            encoded = logo_file.read().strip()
        with tempfile.NamedTemporaryFile(delete=False, suffix=".jpg") as tmp:
            tmp.write(base64.b64decode(encoded))
            return tmp.name
    except Exception:
        return None


def _add_logo(paragraph, width=3.4):
    logo_path = _logo_temp_path()
    if not logo_path:
        return False
    try:
        paragraph.add_run().add_picture(logo_path, width=Inches(width))
        return True
    except Exception:
        return False
    finally:
        try:
            os.remove(logo_path)
        except Exception:
            pass


def _configure(doc, data):
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.68)
    section.right_margin = Inches(0.68)

    header = section.header.paragraphs[0]
    header.text = f"Valhalla Premium Tax Strategy Report | {_safe(data.get('client_name'), 'Client')} | {_safe(data.get('tax_year'), 'Tax Year')}"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _paragraph(header, size=7.5, color=TEXT_GRAY, after=0)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Valhalla Tax & Finance LLC | Confidential planning document | Page ")
    _page_number(footer)
    _paragraph(footer, size=7.3, color=TEXT_GRAY, after=0)

    doc.styles["Normal"].font.name = "Aptos"
    doc.styles["Normal"].font.size = Pt(9.6)


def _section_title(doc, title, kicker=None):
    if kicker:
        k = doc.add_paragraph(kicker.upper())
        _paragraph(k, size=7.3, color=BRAND_GOLD, bold=True, before=5, after=0)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(title)
    r.font.name = "Aptos Display"
    r.font.size = Pt(15.5)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(BRAND_RED)
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "7")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), LINE_GRAY)
    border.append(bottom)
    p._p.get_or_add_pPr().append(border)


def _section_band(doc, title, subtitle=""):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_widths(table, [7.15])
    cell = table.cell(0, 0)
    _format_cell(cell, fill=BRAND_BLACK, color=WHITE, border=BRAND_BLACK)
    p = cell.paragraphs[0]
    p.text = ""
    k = p.add_run(title.upper())
    k.font.name = "Aptos Display"
    k.font.size = Pt(12.2)
    k.font.bold = True
    k.font.color.rgb = RGBColor.from_string(WHITE)
    if subtitle:
        sub = cell.add_paragraph(subtitle)
        _paragraph(sub, size=8.2, color="D7DEE6", after=0)
    doc.add_paragraph().paragraph_format.space_after = Pt(5)


def _divider(doc, color=LINE_GRAY):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_widths(table, [7.15])
    cell = table.cell(0, 0)
    _format_cell(cell, fill=color, border=color)
    _cell_margins(cell, 12, 0, 12, 0)
    cell.paragraphs[0].text = ""
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def _label_value(cell, label, value, value_color=INK, value_size=14.0):
    p = cell.paragraphs[0]
    r = p.add_run(label.upper() + "\n")
    r.font.name = "Aptos"
    r.font.size = Pt(7.1)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(TEXT_GRAY)
    v = p.add_run(str(value))
    v.font.name = "Aptos Display"
    v.font.size = Pt(value_size)
    v.font.bold = True
    v.font.color.rgb = RGBColor.from_string(value_color)


def _callout(doc, title, body, fill=LIGHT_RED, accent=BRAND_RED):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_widths(table, [0.14, 7.12])
    bar = table.cell(0, 0)
    body_cell = table.cell(0, 1)
    _set_shading(bar, accent)
    _set_border(bar, accent, "0")
    _cell_margins(bar, 40, 20, 40, 20)
    _format_cell(body_cell, fill=fill, border="E5D5D8")
    p = body_cell.paragraphs[0]
    p.add_run(title)
    _paragraph(p, size=9.6, color=accent, bold=True, after=2)
    p2 = body_cell.add_paragraph(_limit(body, 560))
    _paragraph(p2, size=8.9, color=INK, after=0)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def _simple_table(doc, rows, widths, header_fill=BRAND_RED, font_size=8.3):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_widths(table, widths)
    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = str(value)
            header = r_idx == 0
            fill = header_fill if header else (WHITE if r_idx % 2 else LIGHT_GRAY)
            color = WHITE if header else INK
            bold = header or c_idx == 0
            align = WD_ALIGN_PARAGRAPH.CENTER if header or c_idx in (0, 1) else WD_ALIGN_PARAGRAPH.LEFT
            _format_cell(cell, fill=fill, color=color, bold=bold, size=font_size, align=align)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return table


def _refund_or_due(data):
    balance = _num(data.get("balance_due", 0))
    refund = _num(data.get("refund", 0))
    if balance > 0:
        return "Balance due", _money(-balance), BRAND_RED
    if refund > 0:
        return "Refund", _money(refund), BRAND_RED
    return "Refund / balance", "$0", BRAND_DARK


def _opportunity_range(data, actions):
    total = 0
    for action in actions or []:
        savings = action.get("estimated_savings") or action.get("tax_impact") or action.get("impact")
        if isinstance(savings, (int, float)):
            total += float(savings)
    if total:
        return _money(total), _money(total * 1.6)
    taxable = _num(data.get("taxable_income", 0))
    sched_c = _num(data.get("schedule_c_net_profit", 0))
    if sched_c > 0:
        return "$2,500", "$8,000+"
    if taxable > 0:
        return "$1,000", "$4,000+"
    return "$500", "$2,500+"


def _tax_efficiency_score(data, actions):
    explicit = data.get("tax_efficiency_score") or data.get("tax_efficiency")
    if isinstance(explicit, dict):
        explicit = explicit.get("score") or explicit.get("overall_score")
    if explicit not in (None, ""):
        try:
            return max(1, min(100, int(float(explicit))))
        except Exception:
            pass

    score = 72
    taxable = _num(data.get("taxable_income", 0))
    total_tax = _num(data.get("total_tax", 0))
    agi = _num(data.get("agi", 0))
    refund = _num(data.get("refund", 0))
    balance = _num(data.get("balance_due", 0))
    sched_c = _num(data.get("schedule_c_net_profit", 0))

    if agi and total_tax / agi < 0.08:
        score += 6
    if taxable and taxable < 100000:
        score += 4
    if balance > 2500:
        score -= 8
    if refund > 4000:
        score -= 5
    if sched_c > 0:
        score -= 4
    if actions:
        score += min(6, len(actions))
    return max(40, min(95, int(score)))


def _score_band(score):
    if score >= 85:
        return "Strong", BRAND_GOLD, "Solid tax position; focus on preserving gains and executing the highest-value actions."
    if score >= 70:
        return "Good with opportunities", BRAND_RED, "Planning value exists; prioritize the actions that improve cash flow, deductions, and long-term flexibility."
    return "Needs attention", BRAND_RED, "Several controllable planning items should be addressed before year-end."


def _default_actions(data):
    actions = []
    balance = _num(data.get("balance_due", 0))
    refund = _num(data.get("refund", 0))
    sched_c = _num(data.get("schedule_c_net_profit", 0))
    taxable = _num(data.get("taxable_income", 0))
    state = str(data.get("state", "")).upper()

    if sched_c > 0:
        actions.append({
            "title": "Fund a self-employed retirement plan",
            "estimated_savings": "Potential annual federal savings from deductible contributions",
            "timeline": "Before plan and contribution deadlines",
            "reason": "Schedule C profit creates one of the cleanest recurring planning opportunities: deductible retirement funding that also builds long-term assets.",
        })
        actions.append({
            "title": "Review Schedule C documentation and entity timing",
            "estimated_savings": "Risk reduction and future payroll-tax planning",
            "timeline": "Next 30 to 90 days",
            "reason": "Business profit should be supported by clean records, substantiated expenses, and an annual review of whether entity changes are warranted.",
        })

    if balance > 0:
        actions.append({
            "title": "Correct withholding or estimated tax payments",
            "estimated_savings": "Cash-flow improvement and penalty prevention",
            "timeline": "Immediately",
            "reason": f"The return reflects a balance due of {_money(balance)}. Adjusting payroll withholding or estimates can prevent repeating the shortfall.",
        })
    elif refund > 2500:
        actions.append({
            "title": "Right-size withholding for better cash flow",
            "estimated_savings": "Cash-flow improvement",
            "timeline": "Next payroll cycle",
            "reason": f"The return reflects a refund of {_money(refund)}. Some excess withholding may be redirected toward planning goals during the year.",
        })

    if taxable > 0:
        actions.append({
            "title": "Model bracket-aware Roth and investment planning",
            "estimated_savings": "Long-term tax flexibility",
            "timeline": "Annual review",
            "reason": "Taxable income creates the need to coordinate Roth conversions, capital gains, retirement contributions, and future income timing.",
        })

    if state == "AZ":
        actions.append({
            "title": "Use Arizona credits and deduction timing intentionally",
            "estimated_savings": "State tax reduction opportunity",
            "timeline": "Before year-end",
            "reason": "Arizona planning can add value through charitable credits, school credits, and careful year-end payment timing.",
        })

    if not actions:
        actions.append({
            "title": "Build a year-round tax planning calendar",
            "estimated_savings": "Planning discipline and risk reduction",
            "timeline": "Next 30 days",
            "reason": "The return should become the starting point for proactive planning, not merely a filing record.",
        })
    return actions[:5]


def _normalize_actions(data, strategy_output, roi_output):
    source = data.get("priority_actions") or strategy_output.get("priority_actions") or []
    actions = []
    for item in source:
        if isinstance(item, str):
            actions.append({"title": item, "estimated_savings": "Planning value", "timeline": "Review", "reason": item})
        elif isinstance(item, dict):
            actions.append({
                "title": item.get("title") or item.get("strategy") or item.get("recommendation") or "Planning action",
                "estimated_savings": item.get("estimated_savings") or item.get("tax_impact") or item.get("impact") or "Planning value",
                "timeline": item.get("timeline") or item.get("timing") or "Review",
                "reason": item.get("reason") or item.get("advisor_reasoning") or item.get("why") or "Client-specific planning opportunity.",
                "score": item.get("score") or item.get("priority_score"),
            })

    for item in data.get("roi_strategies") or roi_output.get("roi_strategies") or []:
        if isinstance(item, dict):
            actions.append({
                "title": item.get("strategy") or item.get("title") or "ROI strategy",
                "estimated_savings": _money(item.get("estimated_savings", 0)) if item.get("estimated_savings") is not None else "Quantified opportunity",
                "timeline": item.get("timeline") or "Review",
                "reason": item.get("advisor_reasoning") or "Quantified planning opportunity identified.",
                "score": item.get("score"),
            })

    if not actions:
        actions = _default_actions(data)
    actions.sort(key=lambda x: _num(x.get("score", 0)), reverse=True)
    return actions[:6]


def _cover(doc, data, actions):
    logo = doc.add_paragraph()
    logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if not _add_logo(logo, width=3.25):
        logo.add_run("VALHALLA TAX & FINANCE LLC")
        _paragraph(logo, size=15, color=BRAND_RED, bold=True, after=3)
    logo.paragraph_format.space_after = Pt(10)

    band = doc.add_table(rows=1, cols=1)
    band.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_widths(band, [7.25])
    cell = band.cell(0, 0)
    _format_cell(cell, fill=BRAND_BLACK, color=WHITE, border=BRAND_BLACK)
    p = cell.paragraphs[0]
    p.add_run("Premium tax strategy | Client advisory deliverable | Implementation roadmap")
    _paragraph(p, size=9.0, color="D7DEE6", bold=True, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)
    title = doc.add_paragraph()
    title.add_run("Tax Strategy\nImplementation Report")
    _paragraph(title, size=29, color=BRAND_RED, bold=True, after=3)

    subtitle = doc.add_paragraph("A premium client-facing roadmap for reducing avoidable tax drag, improving cash-flow control, and turning the tax return into an implementation plan.")
    _paragraph(subtitle, size=10.4, color=TEXT_GRAY, after=10)

    meta = doc.add_table(rows=1, cols=4)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_widths(meta, [2.15, 1.4, 1.55, 2.15])
    score = _tax_efficiency_score(data, actions)
    band_label, band_color, _ = _score_band(score)
    values = [
        ("Client", _safe(data.get("client_name"), "Client"), BRAND_RED),
        ("Tax Year", _safe(data.get("tax_year"), "2025"), INK),
        ("Efficiency", f"{score}/100", band_color),
        ("Prepared By", "Scott Kunz, ChFC, TPCP, EA", INK),
    ]
    for idx, (label, value, color) in enumerate(values):
        meta_cell = meta.cell(0, idx)
        _format_cell(meta_cell, fill=SOFT_SLATE, border="DDE3EA")
        _label_value(meta_cell, label, value, color, 11.6)

    low, high = _opportunity_range(data, actions)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    proof = doc.add_table(rows=1, cols=3)
    proof.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_widths(proof, [2.35, 2.35, 2.35])
    proof_items = [
        ("What this plan does", "Ranks the highest-value planning moves and turns them into a client action sequence."),
        ("What the client sees", "Clear dollars, timing, priority, and why each strategy matters."),
        ("How to use it", "Use this report as the agenda for the implementation meeting."),
    ]
    for idx, (label, body) in enumerate(proof_items):
        proof_cell = proof.cell(0, idx)
        _format_cell(proof_cell, fill=WHITE if idx != 1 else LIGHT_GOLD, border="E2E8F0")
        _add_cell_text(proof_cell, label.upper(), size=7.2, color=DEEP_GOLD, bold=True)
        p_body = proof_cell.add_paragraph(body)
        _paragraph(p_body, size=8.2, color=INK, after=0)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    _callout(
        doc,
        "Executive Value Thesis",
        f"The strategies in this report point to an estimated annual planning opportunity of approximately {low} to {high}, depending on contribution levels, timing, documentation, and final implementation choices.",
        fill=LIGHT_GOLD,
        accent=BRAND_GOLD,
    )

    summary = data.get("advisor_summary") or data.get("planning_summary") or data.get("top_planning_focus") or "This report turns the reviewed return into an implementation plan: what matters most, why it matters, and what the client should act on first."
    _callout(doc, "Advisor Summary", summary, fill=LIGHT_RED, accent=BRAND_RED)

    toc = doc.add_table(rows=6, cols=2)
    toc.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_widths(toc, [0.52, 6.72])
    items = [
        ("01", "Executive opportunity dashboard"),
        ("02", "Confirmed tax position and planning read"),
        ("03", "Client decision matrix"),
        ("04", "Roth conversion and capital gain scenarios"),
        ("05", "Visual tax snapshot"),
        ("06", "Priority strategies and action roadmap"),
    ]
    for row, (num, label) in zip(toc.rows, items):
        row.cells[0].text = num
        row.cells[1].text = label
        _format_cell(row.cells[0], fill=BRAND_RED, color=WHITE, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, border=BRAND_RED)
        _format_cell(row.cells[1], fill=WHITE, color=INK, size=9.1)
    doc.add_page_break()


def _dashboard(doc, data, actions):
    _section_band(doc, "Executive Opportunity Dashboard", "The planning conversation starts here: current position, cash-flow issue, priority range, and top actions.")
    low, high = _opportunity_range(data, actions)
    refund_label, refund_value, refund_color = _refund_or_due(data)
    cards = [
        ("AGI", _money(data.get("agi", 0)), "Phaseout and planning baseline", BRAND_RED),
        ("Taxable Income", _money(data.get("taxable_income", 0)), "Bracket management base", INK),
        ("Total Tax", _money(data.get("total_tax", 0)), "Current federal tax cost", INK),
        (refund_label, refund_value, "Cash-flow planning target", refund_color),
        ("Planning Range", f"{low}-{high}", "Estimated annual opportunity", BRAND_GOLD),
    ]
    table = doc.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_widths(table, [1.42, 1.42, 1.42, 1.42, 1.62])
    for idx, (label, value, note, color) in enumerate(cards):
        cell = table.cell(0, idx)
        _format_cell(cell, fill=LIGHT_GRAY, border="E1E5EA")
        _label_value(cell, label, value, color, 11.4 if idx < 4 else 10.5)
        p = cell.add_paragraph(note)
        _paragraph(p, size=7.2, color=TEXT_GRAY, after=0)

    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    _tax_efficiency_panel(doc, data, actions)
    _strategy_cards(doc, actions[:3], compact=True)
    _divider(doc, "E8EDF3")


def _decision_matrix(doc, data, actions):
    _section_title(doc, "Client Decision Matrix", "What to approve")
    intro = doc.add_paragraph("This page turns the plan into decisions. It is designed to help the client see which strategies deserve approval, what each one affects, and when the work should begin.")
    _paragraph(intro, size=8.9, color=TEXT_GRAY, after=6)

    rows = [["Priority", "Strategy", "Client Value", "Timing", "Decision"]]
    for idx, action in enumerate(actions[:5], start=1):
        rows.append([
            f"#{idx}",
            _safe(action.get("title"), "Planning action"),
            _limit(action.get("estimated_savings", "Planning value"), 90),
            _safe(action.get("timeline"), "Review"),
            "Approve / Model / Defer",
        ])
    _simple_table(doc, rows, [0.72, 2.15, 1.75, 1.25, 1.38], header_fill=BRAND_BLACK, font_size=7.9)

    grid = doc.add_table(rows=1, cols=3)
    grid.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_widths(grid, [2.35, 2.35, 2.35])
    cells = [
        ("Approve", "Move forward when the tax impact, timing, and documentation are clear."),
        ("Model", "Run a scenario when the strategy depends on income, contribution level, or entity timing."),
        ("Defer", "Park the item when the facts are incomplete or the cost outweighs current-year value."),
    ]
    fills = [LIGHT_GOLD, LIGHT_GRAY, LIGHT_RED]
    accents = [DEEP_GOLD, BRAND_BLACK, BRAND_RED]
    for idx, (label, body) in enumerate(cells):
        cell = grid.cell(0, idx)
        _format_cell(cell, fill=fills[idx], border="DDE3EA")
        _add_cell_text(cell, label.upper(), size=8.0, color=accents[idx], bold=True)
        p = cell.add_paragraph(body)
        _paragraph(p, size=8.0, color=INK, after=0)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def _value_roadmap(doc, data, actions):
    _section_title(doc, "Planning Value Roadmap", "How value becomes action")
    rows = [["Stage", "Purpose", "Client-Facing Result"]]
    roadmap = [
        ("1. Diagnose", "Confirm the tax baseline, cash-flow issue, and bracket position.", "Client understands what the return is telling them."),
        ("2. Prioritize", "Rank strategies by impact, urgency, timing, and implementation friction.", "Client sees where the first dollars of effort should go."),
        ("3. Model", "Quantify selected strategies before implementation.", "Client can approve strategies with clearer expectations."),
        ("4. Implement", "Execute approved actions with documentation and deadlines.", "Client receives a plan that turns into measurable work."),
    ]
    rows.extend(roadmap)
    _simple_table(doc, rows, [1.1, 3.05, 3.1], header_fill=BRAND_RED, font_size=8.0)

    _callout(
        doc,
        "Advisor Positioning Note",
        "This report is intentionally structured as an implementation document rather than a tax summary. The goal is to move the client from return review to approved planning decisions.",
        fill=LIGHT_GRAY,
        accent=BRAND_BLACK,
    )


def _tax_efficiency_panel(doc, data, actions):
    score = _tax_efficiency_score(data, actions)
    label, color, read = _score_band(score)
    low, high = _opportunity_range(data, actions)
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_widths(table, [1.3, 5.95])
    score_cell = table.cell(0, 0)
    read_cell = table.cell(0, 1)
    _format_cell(score_cell, fill=color, color=WHITE, bold=True, border=color)
    score_cell.paragraphs[0].text = ""
    p = score_cell.paragraphs[0]
    r = p.add_run(f"{score}\n")
    r.font.name = "Aptos Display"
    r.font.size = Pt(24)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(WHITE)
    label_run = p.add_run("TAX EFFICIENCY")
    label_run.font.name = "Aptos"
    label_run.font.size = Pt(7.4)
    label_run.bold = True
    label_run.font.color.rgb = RGBColor.from_string(WHITE)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)

    _format_cell(read_cell, fill=WHITE, border="DDE3EA")
    h = read_cell.paragraphs[0]
    h.add_run(f"{label} | Priority planning range: {low}-{high}")
    _paragraph(h, size=10.0, color=INK, bold=True, after=2)
    body = read_cell.add_paragraph(read)
    _paragraph(body, size=8.4, color=TEXT_GRAY, after=0)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def _tax_position(doc, data):
    _section_title(doc, "Confirmed Tax Position", "Return facts")
    rows = [
        ["Tax Fact", "Confirmed Amount", "Planning Read"],
        ["Filing status", _safe(data.get("filing_status"), "Not provided"), "Controls brackets, standard deduction, credits, and phaseouts"],
        ["Dependents", _safe(data.get("dependents"), "0"), "Important for credits, family planning, education, and payroll strategies"],
        ["AGI", _money(data.get("agi", 0)), "Baseline for phaseouts, credits, IRMAA, and state planning"],
        ["Taxable income", _money(data.get("taxable_income", 0)), "Primary driver for bracket-aware retirement and investment strategy"],
        ["Total tax", _money(data.get("total_tax", 0)), "Current tax-cost benchmark"],
        [_refund_or_due(data)[0], _refund_or_due(data)[1], "Withholding and estimated payments should be tuned before year-end"],
    ]
    _simple_table(doc, rows, [1.52, 1.78, 3.95], font_size=8.4)

    focus = data.get("top_planning_focus") or "The return should be used as a planning baseline for retirement funding, withholding, deductions, and state-specific opportunities."
    _callout(doc, "Advisor Read", focus, fill=LIGHT_GRAY, accent=BRAND_BLACK)


def _business_section(doc, data):
    gross = _num(data.get("schedule_c_gross_revenue", 0))
    net = _num(data.get("schedule_c_net_profit", 0))
    if not gross and not net:
        return
    expenses = max(gross - net, 0)
    _section_title(doc, "Business Owner Planning", "Schedule C")
    rows = [
        ["Metric", "Amount", "Planning Interpretation"],
        ["Gross revenue", _money(gross), "Business activity is large enough to justify proactive tax structure and documentation review"],
        ["Expenses", _money(expenses), f"Expense ratio is {_pct(expenses, gross)}; confirm substantiation and business purpose"],
        ["Net profit", _money(net), f"Net margin is {_pct(net, gross)}; retirement funding and entity review should be modeled"],
        ["Contract labor", _money(data.get("contract_labor", 0)), "Confirm worker classification, W-9 files, and 1099 documentation"],
    ]
    _simple_table(doc, rows, [1.55, 1.45, 4.25], font_size=8.3)
    _callout(doc, "Business Planning Priority", "Tie Schedule C profit to clean books, documented expenses, retirement funding, quarterly tax planning, and an annual entity trigger review.", fill=LIGHT_GOLD, accent=BRAND_GOLD)


def _roth_conversion_section(doc, data):
    taxable = _num(data.get("taxable_income", 0))
    if taxable <= 0:
        return
    status = data.get("filing_status", "single")
    explicit = data.get("roth_conversion_amount") or data.get("recommended_roth_conversion")
    if explicit not in (None, ""):
        target = max(0, _num(explicit))
    else:
        rate = _marginal_rate(taxable, status)
        if rate <= 0.12:
            target = 60000
        elif rate <= 0.22:
            target = 50000
        elif rate <= 0.24:
            target = 30000
        else:
            target = 20000
    if target <= 0:
        return

    _section_title(doc, "Roth Conversion Tax Cost Scenarios", "1, 3, and 5 year view")
    note = doc.add_paragraph("The table below illustrates federal tax cost if the same total Roth conversion target is completed all at once or spread across multiple years. It is a report illustration only; final conversion sizing should be modeled with actual current-year income, state tax, IRMAA exposure, and cash available to pay the tax.")
    _paragraph(note, size=8.4, color=TEXT_GRAY, after=6)

    rows = [["Scenario", "Annual Conversion", "Estimated Annual Federal Tax", "Total Federal Tax", "Effective Rate"]]
    for years in (1, 3, 5):
        annual = target / years
        annual_tax = max(0, _ordinary_tax(taxable + annual, status) - _ordinary_tax(taxable, status))
        total_tax = annual_tax * years
        rows.append([
            f"{years} year" if years == 1 else f"{years} years",
            _money(annual),
            _money(annual_tax),
            _money(total_tax),
            _rate(total_tax / target if target else 0),
        ])
    _simple_table(doc, rows, [1.0, 1.55, 1.85, 1.45, 1.4], header_fill=BRAND_BLACK, font_size=8.0)

    _chart(
        doc,
        "Roth Conversion Federal Tax Cost by Timing",
        ["1 year", "3 years", "5 years"],
        [max(0, _ordinary_tax(taxable + target, status) - _ordinary_tax(taxable, status)),
         max(0, _ordinary_tax(taxable + target / 3, status) - _ordinary_tax(taxable, status)) * 3,
         max(0, _ordinary_tax(taxable + target / 5, status) - _ordinary_tax(taxable, status)) * 5],
        "Compares estimated federal tax cost for the same conversion target under different implementation timelines. Lower bars indicate lower projected federal tax cost under this simplified illustration.",
        kind="bar",
        width=5.65,
    )

    _callout(
        doc,
        "Roth Planning Read",
        "A multi-year conversion schedule can smooth tax cost and may preserve bracket flexibility. The preferred option is usually the one that converts enough to improve long-term tax diversification without pushing the client into avoidable higher-rate or premium-surcharge territory.",
        fill=LIGHT_GOLD,
        accent=BRAND_GOLD,
    )


def _capital_gain_section(doc, data):
    gains = _num(data.get("capital_gains", data.get("capital_gain", 0)))
    dividends = _num(data.get("qualified_dividends", data.get("dividend_income", 0)))
    losses = abs(min(0, gains))
    preferred = max(0, gains) + max(0, dividends)
    taxable = _num(data.get("taxable_income", 0))
    if not preferred and not losses and not taxable:
        return

    status = data.get("filing_status", "single")
    ordinary_base = max(0, taxable - preferred)
    zero_top, fifteen_top = _ltcg_thresholds(status)
    zero_room = max(0, zero_top - ordinary_base)
    fifteen_room = max(0, fifteen_top - max(ordinary_base, zero_top))
    at_zero, at_fifteen, at_twenty, preferred_tax = _estimate_ltcg_tax(ordinary_base, preferred, status)

    _section_title(doc, "Capital Gain and Loss Planning", "Investment tax review")
    rows = [
        ["Planning Item", "Amount", "Client Meaning"],
        ["Long-term gains reported", _money(gains), "Use gain harvesting or loss harvesting intentionally rather than reactively"],
        ["Dividend income", _money(dividends), "Confirm qualified vs. ordinary dividend character before final tax modeling"],
        ["Estimated ordinary taxable base", _money(ordinary_base), "Preferred income is stacked on top of ordinary income for capital gain bracket purposes"],
        ["0% LTCG bracket room", _money(zero_room), "Potential room where qualified dividends or long-term gains may be taxed at 0% federally"],
        ["15% LTCG bracket room", _money(fifteen_room), "Room before the federal 20% long-term capital gain bracket begins"],
        ["Estimated federal tax on preferred income", _money(preferred_tax), "Approximate federal tax on reported gains and dividends in this illustration"],
    ]
    _simple_table(doc, rows, [1.85, 1.35, 4.05], header_fill=BRAND_BLACK, font_size=8.0)

    bracket_rows = [
        ["Capital Gain Bracket", "Income Assigned", "Estimated Tax"],
        ["0% bracket", _money(at_zero), "$0"],
        ["15% bracket", _money(at_fifteen), _money(at_fifteen * 0.15)],
        ["20% bracket", _money(at_twenty), _money(at_twenty * 0.20)],
        ["Total preferred income", _money(preferred), _money(preferred_tax)],
    ]
    _simple_table(doc, bracket_rows, [1.8, 1.75, 1.55], header_fill=BRAND_RED, font_size=8.0)

    _chart(
        doc,
        "Preferred Income by Capital Gain Bracket",
        ["0% LTCG", "15% LTCG", "20% LTCG"],
        [at_zero, at_fifteen, at_twenty],
        "Shows where reported long-term gains and dividend income fall after ordinary taxable income is considered. This is a bracket-stacking view, not a strategy ranking chart.",
        kind="bar",
        width=5.65,
    )

    _callout(
        doc,
        "Capital Gain Planning Read",
        "If taxable investments exist, review unrealized gains and losses before year-end. The planning goal is to coordinate gain harvesting, loss harvesting, charitable gifting of appreciated assets, and bracket management before trades are placed.",
        fill=LIGHT_RED if gains > 0 else LIGHT_GOLD,
        accent=BRAND_RED if gains > 0 else BRAND_GOLD,
    )


def _create_chart(path, title, labels, values, kind="barh"):
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        import matplotlib.ticker as mtick

        values = [_num(v) for v in values]
        colors = ["#981E26", "#B7892B", "#242424", "#9CA3AF", "#D8DEE5", "#5C6670"]
        plt.rcParams["font.family"] = "DejaVu Sans"
        fig, ax = plt.subplots(figsize=(6.2, 2.25))
        fig.patch.set_facecolor("white")
        ax.set_facecolor("white")

        if kind == "donut":
            safe = [max(v, 0) for v in values]
            total = sum(safe) or 1
            ax.pie(safe, labels=labels, startangle=90, colors=colors[: len(labels)], wedgeprops={"width": 0.45, "edgecolor": "white"}, textprops={"fontsize": 7.2})
            ax.text(0, 0, _money(total), ha="center", va="center", fontsize=10, fontweight="bold", color="#981E26")
        elif kind == "bar":
            ax.bar(labels, values, color=colors[: len(labels)], width=0.56)
            ax.yaxis.set_major_formatter(mtick.StrMethodFormatter("${x:,.0f}"))
            ax.tick_params(axis="x", labelsize=7.2, rotation=18, colors="#1F2933")
            ax.tick_params(axis="y", labelsize=7.4, colors="#5C6670")
            ax.grid(axis="y", color="#E6E9ED", linewidth=0.8)
        else:
            order = list(range(len(labels)))[::-1]
            ax.barh([labels[i] for i in order], [values[i] for i in order], color=[colors[i % len(colors)] for i in order], height=0.54)
            ax.xaxis.set_major_formatter(mtick.StrMethodFormatter("${x:,.0f}"))
            ax.tick_params(axis="x", labelsize=7.3, colors="#5C6670")
            ax.tick_params(axis="y", labelsize=8.0, colors="#1F2933")
            ax.grid(axis="x", color="#E6E9ED", linewidth=0.8)

        ax.set_title(title, fontsize=10.2, fontweight="bold", color="#242424", pad=8)
        for spine in ax.spines.values():
            spine.set_visible(False)
        fig.tight_layout(pad=1.0)
        fig.savefig(path, dpi=175, bbox_inches="tight")
        plt.close(fig)
        return True
    except Exception:
        return False


def _chart(doc, title, labels, values, caption, kind="barh", width=5.95):
    with tempfile.NamedTemporaryFile(delete=False, suffix=".png") as tmp:
        chart_path = tmp.name
    try:
        if _create_chart(chart_path, title, labels, values, kind=kind):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(4)
            p.add_run().add_picture(chart_path, width=Inches(width))
            cap = doc.add_paragraph(caption)
            _paragraph(cap, size=7.8, color=TEXT_GRAY, italic=True, after=7, align=WD_ALIGN_PARAGRAPH.CENTER)
        else:
            _chart_fallback_table(doc, title, labels, values, caption)
    finally:
        try:
            os.remove(chart_path)
        except Exception:
            pass


def _chart_fallback_table(doc, title, labels, values, caption):
    heading = doc.add_paragraph(title)
    _paragraph(heading, size=10.0, color=BRAND_BLACK, bold=True, after=4, align=WD_ALIGN_PARAGRAPH.CENTER)
    numeric_values = [_num(v) for v in values]
    max_value = max([abs(v) for v in numeric_values] or [1]) or 1
    rows = [["Item", "Amount", "Relative View"]]
    for label, value in zip(labels, numeric_values):
        blocks = max(1, int(abs(value) / max_value * 12)) if value else 0
        bar = ("|" * blocks) if blocks else "-"
        rows.append([label, _money(value), bar])
    _simple_table(doc, rows, [2.25, 1.55, 3.1], header_fill=BRAND_BLACK, font_size=8.0)
    cap = doc.add_paragraph(caption)
    _paragraph(cap, size=7.8, color=TEXT_GRAY, italic=True, after=7, align=WD_ALIGN_PARAGRAPH.CENTER)


def _visuals(doc, data, actions):
    _section_band(doc, "Tax Return Snapshot Charts", "These charts summarize return facts only. Strategy tax costs are shown separately in the Roth and capital gain sections.")
    guide_rows = [
        ["Chart", "What It Means"],
        ["Tax return dollar snapshot", "Compares AGI, taxable income, total tax, withholding, and refund or balance due from the return."],
        ["Schedule C economics", "Shows gross revenue, estimated expenses, and net profit or loss when business activity exists."],
    ]
    _simple_table(doc, guide_rows, [2.15, 5.1], header_fill=BRAND_BLACK, font_size=8.0)
    _chart(
        doc,
        "Tax Return Dollar Snapshot",
        ["AGI", "Taxable Income", "Total Tax", "Withholding", _refund_or_due(data)[0]],
        [data.get("agi", 0), data.get("taxable_income", 0), data.get("total_tax", 0), data.get("federal_withholding", 0), abs(_num(data.get("balance_due", 0)) or _num(data.get("refund", 0)))],
        "This chart is not a savings projection. It simply shows the major dollar amounts from the return so the client can see the planning baseline.",
        kind="barh",
        width=5.75,
    )

    gross = _num(data.get("schedule_c_gross_revenue", 0))
    net = _num(data.get("schedule_c_net_profit", 0))
    if gross or net:
        _chart(
            doc,
            "Schedule C Economics",
            ["Gross Revenue", "Expenses", "Net Profit"],
            [gross, max(gross - net, 0), net],
            "This chart explains the business income picture: gross revenue, estimated deductions, and the resulting net profit or loss.",
            kind="barh",
            width=5.75,
        )


def _strategy_cards(doc, actions, compact=False):
    if not compact:
        _section_band(doc, "Priority Strategy Recommendations", "Each recommendation is presented as an action card with impact, timing, and advisor rationale.")
    rows_per = 1 if compact else min(len(actions), 6)
    for idx, action in enumerate(actions[: rows_per if compact else 6], start=1):
        table = doc.add_table(rows=1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        _set_widths(table, [0.56, 6.7])
        num = table.cell(0, 0)
        body = table.cell(0, 1)
        num.text = f"{idx}"
        _format_cell(num, fill=BRAND_RED, color=WHITE, bold=True, size=13.5, align=WD_ALIGN_PARAGRAPH.CENTER, border=BRAND_RED)
        _format_cell(body, fill=WHITE, border="DDE3EA")
        title = body.paragraphs[0]
        title.add_run(_safe(action.get("title"), "Planning action"))
        _paragraph(title, size=10.5, color=INK, bold=True, after=2)
        details = body.add_paragraph()
        score = action.get("score")
        if score not in (None, ""):
            details.add_run("Score: ").bold = True
            details.add_run(f"{int(_num(score))}/100   ")
        details.add_run("Impact: ").bold = True
        details.add_run(str(action.get("estimated_savings", "Planning value")))
        details.add_run("   Timing: ").bold = True
        details.add_run(str(action.get("timeline", "Review")))
        _paragraph(details, size=8.4, color=TEXT_GRAY, after=2)
        reason = body.add_paragraph(_limit(action.get("reason", "Client-specific planning opportunity."), 430 if compact else 620))
        _paragraph(reason, size=8.7, color=INK, after=0)
        if not compact:
            decision = body.add_paragraph()
            decision.add_run("Client decision: ").bold = True
            decision.add_run("Approve, model, or defer after advisor review.")
            _paragraph(decision, size=8.1, color=DEEP_GOLD, after=0)
        doc.add_paragraph().paragraph_format.space_after = Pt(2)


def _implementation(doc, actions):
    _section_band(doc, "Implementation Roadmap", "The plan should leave the meeting with owners, timing, and next steps.")
    first = actions[0].get("title", "highest-value planning item") if actions else "highest-value planning item"
    rows = [
        ["Timing", "Client Action", "Advisor Follow-Up"],
        ["Next 30 days", f"Start the first priority: {first}.", "Confirm supporting facts, deadlines, and expected tax impact."],
        ["Next 90 days", "Model retirement, withholding, business, and state planning decisions.", "Prepare implementation projections and client decision points."],
        ["Before year-end", "Execute approved strategies and document the file.", "Review payroll withholding, estimates, charitable credits, and contribution deadlines."],
        ["Annual review", "Update the plan after the next return and income changes.", "Keep the strategy recurring instead of one-time."],
    ]
    _simple_table(doc, rows, [1.05, 3.25, 3.0], font_size=8.1)


def _final_pages(doc, data):
    _section_band(doc, "Final Advisor Recommendation", "The closing page gives the client a clear next move and keeps the plan actionable.")
    final = data.get("final_recommendation") or data.get("planning_summary") or "The strongest planning value comes from prioritizing the right strategies in the right order. The client should focus first on actions that improve tax efficiency, reduce compliance risk, improve cash-flow predictability, and support long-term wealth building."
    p = doc.add_paragraph(_limit(final, 760))
    _paragraph(p, size=9.3, color=INK, after=8)

    _callout(doc, "Client Next Step", "Review the priority actions, select the strategies to implement, and schedule a follow-up planning meeting before the next major tax deadline.", fill=LIGHT_GOLD, accent=BRAND_GOLD)

    _section_title(doc, "Prepared By", "Advisor contact")
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_widths(table, [3.35, 3.9])
    left = table.cell(0, 0)
    right = table.cell(0, 1)
    left.text = "Scott Kunz, ChFC, TPCP\nEnrolled Agent and Financial Advisor\nValhalla Tax & Finance LLC"
    right.text = "7055 W Bell Rd, Suite B20, Glendale, AZ 85308\n(623) 887-7921\nskunz@valhallataxservice.com\nwww.valhallataxservice.com"
    _format_cell(left, fill=LIGHT_RED, bold=True, size=8.7)
    _format_cell(right, fill=WHITE, size=8.5)

    _section_title(doc, "Important Planning Notes", "Disclosure")
    disc = doc.add_paragraph("The savings estimates in this report are planning illustrations, not guaranteed outcomes. Actual results depend on final income, filing status, business-use percentages, payroll requirements, documentation, entity costs, state law, and implementation timing. Strategies should be implemented with proper documentation and professional review.")
    _paragraph(disc, size=8.1, color=TEXT_GRAY, after=0)


def _build_strategy_output(data):
    if data.get("priority_actions") or data.get("dynamic_sections"):
        return {"priority_actions": data.get("priority_actions", []), "dynamic_sections": data.get("dynamic_sections", [])}
    if generate_dynamic_tax_strategy is None:
        return {"priority_actions": _default_actions(data), "dynamic_sections": []}
    try:
        result = generate_dynamic_tax_strategy(data)
        if not result.get("priority_actions"):
            result["priority_actions"] = _default_actions(data)
        return result
    except Exception:
        return {"priority_actions": _default_actions(data), "dynamic_sections": []}


def _build_roi_output(data):
    if data.get("roi_strategies"):
        return {"tax_efficiency": {}, "roi_strategies": data.get("roi_strategies", [])}
    if generate_roi_analysis is None:
        return {"tax_efficiency": {}, "roi_strategies": []}
    try:
        return generate_roi_analysis(data)
    except Exception:
        return {"tax_efficiency": {}, "roi_strategies": []}


def _dynamic_sections(doc, sections):
    if not sections:
        return
    _section_title(doc, "Client-Specific Planning Detail", "Strategy modules")
    for item in sections[:6]:
        if isinstance(item, dict):
            heading = item.get("section") or item.get("title") or "Planning Module"
            body = item.get("body") or item.get("summary") or ""
        else:
            heading = "Planning Module"
            body = str(item)
        h = doc.add_paragraph(heading)
        _paragraph(h, size=10.0, color=BRAND_RED, bold=True, after=2)
        b = doc.add_paragraph(_limit(body, 650))
        _paragraph(b, size=8.8, color=INK, after=5)


def generate_valhalla_docx_report(data: dict, output_path: str = "valhalla_premium_report.docx"):
    data = data or {}
    strategy_output = _build_strategy_output(data)
    roi_output = _build_roi_output(data)
    actions = _normalize_actions(data, strategy_output, roi_output)
    dynamic_sections = strategy_output.get("dynamic_sections", [])

    doc = Document()
    _configure(doc, data)

    _cover(doc, data, actions)
    _dashboard(doc, data, actions)
    _tax_position(doc, data)
    _decision_matrix(doc, data, actions)
    _business_section(doc, data)
    _roth_conversion_section(doc, data)
    _capital_gain_section(doc, data)
    _visuals(doc, data, actions)
    _value_roadmap(doc, data, actions)
    doc.add_page_break()
    _strategy_cards(doc, actions, compact=False)
    _dynamic_sections(doc, dynamic_sections)
    _implementation(doc, actions)
    doc.add_page_break()
    _final_pages(doc, data)

    doc.save(output_path)
    return output_path
